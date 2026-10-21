from pathlib import Path
import openpyxl
import pytest

from app.application.caso_uso_importar_carreras import ErrorImportacionCarreras
from app.domain.carrera import Carrera


def test_importacion_archivo_real_y_comparacion_filas(
    caso_uso_importar, repositorio_prueba, ruta_excel_real: Path
):
    """Verifica la carga del archivo Excel ampliado y compara cada registro con la fila original."""
    assert ruta_excel_real.exists(), f"No existe el archivo Excel en {ruta_excel_real}"

    resultado = caso_uso_importar.ejecutar(ruta_excel_real)
    assert resultado.total_leidos == 35, f"Se esperaban 35 registros leídos, se obtuvieron {resultado.total_leidos}"
    assert resultado.total_insertados == 35
    assert resultado.total_actualizados == 0
    assert repositorio_prueba.contar() == 35

    # Lectura directa del archivo Excel para contrastar fila por fila
    libro = openpyxl.load_workbook(str(ruta_excel_real), data_only=True)
    hoja = libro["Carreras"]
    encabezados = [c.value for c in hoja[1]]
    idx_id = encabezados.index("id_carrera") + 1
    idx_carrera = encabezados.index("carrera") + 1
    idx_institucion = encabezados.index("institucion") + 1
    idx_duracion = encabezados.index("duracion_meses") + 1
    idx_modalidad = encabezados.index("modalidad") + 1
    idx_ubicacion = encabezados.index("ubicacion") + 1
    idx_tipo = encabezados.index("tipo_universidad") + 1

    for r in range(2, hoja.max_row + 1):
        id_excel = str(hoja.cell(row=r, column=idx_id).value).strip()
        carrera_db = repositorio_prueba.obtener_por_id(id_excel)
        assert carrera_db is not None, f"No se encontró en base de datos la carrera {id_excel}"

        assert carrera_db.carrera == str(hoja.cell(row=r, column=idx_carrera).value).strip()
        assert carrera_db.institucion == str(hoja.cell(row=r, column=idx_institucion).value).strip()
        assert carrera_db.duracion_meses == int(hoja.cell(row=r, column=idx_duracion).value)
        assert carrera_db.modalidad == str(hoja.cell(row=r, column=idx_modalidad).value).strip()
        assert carrera_db.ubicacion == str(hoja.cell(row=r, column=idx_ubicacion).value).strip()
        assert carrera_db.tipo_universidad == str(hoja.cell(row=r, column=idx_tipo).value).strip()


def test_idempotencia_importacion(caso_uso_importar, repositorio_prueba, ruta_excel_real: Path):
    """Comprueba que ejecutar la carga repetidas veces no duplica carreras ni altera la consistencia."""
    # Primera ejecución: inserción de las 35 carreras
    res1 = caso_uso_importar.ejecutar(ruta_excel_real)
    assert res1.total_leidos == 35
    assert res1.total_insertados == 35
    assert res1.total_actualizados == 0
    assert repositorio_prueba.contar() == 35

    # Segunda ejecución: actualización de las 35 carreras (0 insertadas)
    res2 = caso_uso_importar.ejecutar(ruta_excel_real)
    assert res2.total_leidos == 35
    assert res2.total_insertados == 0
    assert res2.total_actualizados == 35
    assert repositorio_prueba.contar() == 35

    # Tercera ejecución: confirmación de estabilidad
    res3 = caso_uso_importar.ejecutar(ruta_excel_real)
    assert res3.total_leidos == 35
    assert res3.total_insertados == 0
    assert res3.total_actualizados == 35
    assert repositorio_prueba.contar() == 35


def test_conservacion_rangos_y_costos(caso_uso_importar, repositorio_prueba, ruta_excel_real: Path):
    """Verifica que los rangos institucionales se conserven sin inventar precios fijos y que los costos se guarden adecuadamente."""
    caso_uso_importar.ejecutar(ruta_excel_real)

    # 1. Rango institucional UPN (C29 - C35)
    for c_id in [f"C{i:02d}" for i in range(29, 36)]:
        upn = repositorio_prueba.obtener_por_id(c_id)
        assert upn is not None, f"Carrera {c_id} no encontrada"
        assert upn.costo_es_rango is True
        assert upn.costo_minimo == 460.0
        assert upn.costo_maximo == 2605.0
        # Regla clave: no inventar precio fijo a partir de rango institucional
        assert upn.costo_monto is None
        assert "460" in upn.costo_texto and "2605" in upn.costo_texto
        assert "Rango institucional UPN" in upn.alcance_costo
        assert "Por cuota" in upn.periodicidad_costo

    # 2. Costo fijo aportado UTP (ej: C01)
    utp = repositorio_prueba.obtener_por_id("C01")
    assert utp is not None
    assert utp.costo_es_rango is False
    assert utp.costo_minimo == 775.0
    assert utp.costo_maximo == 775.0
    assert utp.costo_monto == 775.0
    assert utp.costo_texto == "775"
    assert utp.periodicidad_costo == "Mensual"

    # 3. Gratuidad en universidades públicas (UNI: C20, UNMSM: C05)
    for pub_id in ["C05", "C20"]:
        pub = repositorio_prueba.obtener_por_id(pub_id)
        assert pub is not None
        assert pub.tipo_universidad == "Pública"
        assert pub.costo_es_rango is False
        assert pub.costo_minimo == 0.0
        assert pub.costo_maximo == 0.0
        assert pub.costo_monto == 0.0
        assert pub.costo_texto == "0"
        assert "Gratuidad de enseñanza" in pub.alcance_costo


def test_consultas_y_filtros_repositorio(caso_uso_importar, repositorio_prueba, ruta_excel_real: Path):
    """Comprueba las capacidades de búsqueda, filtrado y conteo del repositorio."""
    caso_uso_importar.ejecutar(ruta_excel_real)

    # Conteo por institución
    assert repositorio_prueba.contar() == 35
    assert repositorio_prueba.contar(institucion="UTP") == 13
    assert repositorio_prueba.contar(institucion="UNMSM") == 6
    assert repositorio_prueba.contar(institucion="UNI") == 9
    assert repositorio_prueba.contar(institucion="UPN") == 7

    # Listado con filtros
    carreras_uni = repositorio_prueba.listar(institucion="UNI")
    assert len(carreras_uni) == 9
    assert all(c.institucion == "UNI" for c in carreras_uni)

    carreras_ingenieria = repositorio_prueba.listar(area_formacion="Ingeniería")
    assert len(carreras_ingenieria) > 0
    assert all(c.area_formacion == "Ingeniería" for c in carreras_ingenieria)

    # Paginación
    primeras_diez = repositorio_prueba.listar(limite=10, desplazamiento=0)
    assert len(primeras_diez) == 10
    assert primeras_diez[0].id_carrera == "C01"

    siguientes = repositorio_prueba.listar(limite=10, desplazamiento=10)
    assert len(siguientes) == 10
    assert siguientes[0].id_carrera == "C11"


def test_validacion_entidad_dominio():
    """Valida que los campos obligatorios del negocio arrojen error cuando no son válidos."""
    carrera = Carrera(
        id_carrera="",
        carrera="Ingeniería",
        institucion="UTP",
        duracion_meses=60,
        modalidad="Presencial",
        ubicacion="Lima Norte",
        costo_texto="775",
        requisitos="Secundaria",
        campo_laboral="TI",
        fuente_carrera="http://",
        fuente_costo="http://",
        fuente_requisitos="http://",
        fecha_revision="2026-10-13",
        fuente_duracion="http://",
        area_formacion="Ingeniería",
        grupo_carrera="sistemas",
        fuente_area="http://",
        criterio_area="Facultad",
        tipo_universidad="Privada",
        ambito_geografico="Lima Norte",
        periodicidad_costo="Mensual",
        alcance_costo="Aportado",
        fuente_ubicacion="http://",
    )
    with pytest.raises(ValueError, match="El identificador de la carrera"):
        carrera.validar()

    carrera.id_carrera = "C01"
    carrera.duracion_meses = -5
    with pytest.raises(ValueError, match="duración en meses"):
        carrera.validar()


def test_error_archivo_inexistente(caso_uso_importar):
    """Comprueba el manejo de error cuando no existe el archivo Excel."""
    with pytest.raises(ErrorImportacionCarreras, match="No se encontró el archivo Excel"):
        caso_uso_importar.ejecutar("ruta/que/no/existe.xlsx")


def test_endpoints_api(cliente_api, caso_uso_importar, ruta_excel_real: Path):
    """Comprueba los endpoints FastAPI de listado, filtros y detalle de carreras."""
    # Importar los datos
    caso_uso_importar.ejecutar(ruta_excel_real)

    # 1. Health check
    resp_health = cliente_api.get("/health")
    assert resp_health.status_code == 200
    assert resp_health.json()["status"] == "ok"

    # 2. Listado general
    resp_listado = cliente_api.get("/api/carreras")
    assert resp_listado.status_code == 200
    datos = resp_listado.json()
    assert datos["total"] == 35
    assert len(datos["elementos"]) == 35

    # 3. Filtro por institución
    resp_filtro = cliente_api.get("/api/carreras?institucion=UNI")
    assert resp_filtro.status_code == 200
    datos_uni = resp_filtro.json()
    assert datos_uni["total"] == 9
    assert all(el["institucion"] == "UNI" for el in datos_uni["elementos"])

    # 4. Detalle de carrera con rango (UPN C29)
    resp_detalle_upn = cliente_api.get("/api/carreras/C29")
    assert resp_detalle_upn.status_code == 200
    c29 = resp_detalle_upn.json()
    assert c29["id_carrera"] == "C29"
    assert c29["costo_es_rango"] is True
    assert c29["costo_minimo"] == 460.0
    assert c29["costo_maximo"] == 2605.0
    assert c29["costo_monto"] is None

    # 4.1. Normalización flexible de identificadores (ej: '2', '02', 'c02' -> 'C02')
    for variante in ["2", "02", "c02", "C02"]:
        resp_var = cliente_api.get(f"/api/carreras/{variante}")
        assert resp_var.status_code == 200
        assert resp_var.json()["id_carrera"] == "C02"

    # 5. Detalle de carrera inexistente
    resp_404 = cliente_api.get("/api/carreras/C999")
    assert resp_404.status_code == 404
    assert "no encontrada" in resp_404.json()["detail"]


def test_filtro_ubicacion_api(cliente_api, caso_uso_importar, ruta_excel_real: Path):
    """Verifica el filtro por ubicacion en el endpoint de listado."""
    caso_uso_importar.ejecutar(ruta_excel_real)

    resp = cliente_api.get("/api/carreras?ubicacion=Virtual")
    assert resp.status_code == 200
    datos = resp.json()
    assert datos["total"] == 2
    assert len(datos["elementos"]) == 2
    assert all(el["ubicacion"] == "Virtual" for el in datos["elementos"])

    resp2 = cliente_api.get("/api/carreras?ubicacion=Lima")
    assert resp2.status_code == 200
    datos2 = resp2.json()
    assert datos2["total"] == 6
    assert all(el["ubicacion"] == "Lima" for el in datos2["elementos"])


def test_paginacion_api(cliente_api, caso_uso_importar, ruta_excel_real: Path):
    """Verifica la paginacion y el total de registros en la respuesta."""
    caso_uso_importar.ejecutar(ruta_excel_real)

    resp = cliente_api.get("/api/carreras?limite=5&desplazamiento=0")
    assert resp.status_code == 200
    datos = resp.json()
    assert datos["total"] == 35
    assert datos["limite"] == 5
    assert datos["desplazamiento"] == 0
    assert len(datos["elementos"]) == 5

    resp2 = cliente_api.get("/api/carreras?limite=5&desplazamiento=30")
    assert resp2.status_code == 200
    datos2 = resp2.json()
    assert len(datos2["elementos"]) == 5
